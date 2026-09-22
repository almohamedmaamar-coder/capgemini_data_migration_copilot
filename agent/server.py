"""Serveur FastAPI de l'agent qualité Profit Center."""
from __future__ import annotations

import io
import os
import uuid

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import BaseModel

load_dotenv()

from detect import detect_workbook  # noqa: E402
from graph import get_graph, llm_configured, model_name, provider_name, run_autonomous_audit  # noqa: E402
from rules_engine import analyze_workbook  # noqa: E402
from tools import store_session  # noqa: E402

app = FastAPI(title="Migration Copilot — agent qualité Profit Center")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ChatIn(BaseModel):
    session_id: str
    message: str


RULE_TO_FILTER = {"R1": "mandatory", "R2": "consistency", "R3": "duplicate"}
FILTER_TO_LABEL = {
    "mandatory": "Voir R1",
    "consistency": "Voir R2",
    "duplicate": "Voir les doublons",
}


@app.get("/health")
def health() -> dict:
    return {
        "status": "ok",
        "llm_configured": llm_configured(),
        "provider": provider_name(),
        "model": model_name(),
    }


@app.post("/analyze")
async def analyze(file: UploadFile = File(...)) -> dict:
    try:
        data = await file.read()
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
        sheets = {name: wb[name] for name in wb.sheetnames}
        models, links = detect_workbook(sheets)
        if not any(not m.ignored for m in models):
            raise ValueError("Aucun tableau détecté dans ce classeur.")
        report = analyze_workbook(models, links, file.filename or "classeur.xlsx")
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:  # noqa: BLE001 — message FR Bubble-up
        raise HTTPException(status_code=400, detail=f"Fichier illisible : {e}")
    session_id = uuid.uuid4().hex[:12]
    store_session(session_id, report)

    # Lancement du raisonnement autonome de l'agent LangGraph
    if llm_configured():
        try:
            synthesis = run_autonomous_audit(session_id)
            report["agentSynthesis"] = synthesis
        except Exception:
            pass

    return {"session_id": session_id, "report": report}


@app.post("/chat")
async def chat(body: ChatIn) -> dict:
    if not llm_configured():
        raise HTTPException(
            status_code=503,
            detail="Modèle IA non configuré — renseignez la clé API dans agent/.env.",
        )
    try:
        result = get_graph().invoke(
            {"messages": [("user", f"[session: {body.session_id}] {body.message}")]},
            config={"configurable": {"thread_id": body.session_id}},
        )
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Erreur modèle : {e}")

    reply = ""
    action_filter = None
    for msg in reversed(result["messages"]):
        tool_calls = getattr(msg, "tool_calls", None) or []
        for call in tool_calls:
            if call.get("name") == "list_issues":
                rule = (call.get("args") or {}).get("rule", "all")
                if rule in RULE_TO_FILTER:
                    action_filter = RULE_TO_FILTER[rule]
        if getattr(msg, "type", "") == "ai" and getattr(msg, "content", None):
            content = msg.content
            if isinstance(content, str) and content.strip():
                reply = content.strip()
                break
    if not reply:
        reply = "Je n'ai pas pu formuler de réponse — reformulez votre question."
    out: dict = {"reply": reply}
    if action_filter:
        out["actionFilter"] = action_filter
        out["actionLabel"] = FILTER_TO_LABEL[action_filter]
    return out


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=int(os.getenv("PORT", "8001")))
